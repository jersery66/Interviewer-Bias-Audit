"""Independent C5 material preparation and R-block sensitivity analysis.

This module is deliberately separate from the locked formal explanation
analysis.  It reads frozen inputs and writes only below
``analysis_v2/16_interviewer_signal_revision`` (or an explicitly supplied
result directory).  C5 review material is split between an owner-only model
crosswalk and blinded reviewer packets with blank human fields.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import subprocess
import warnings
import zipfile
from pathlib import Path
from collections.abc import Mapping, Sequence

import numpy as np
import pandas as pd
from sklearn.exceptions import UndefinedMetricWarning
from sklearn.metrics import cohen_kappa_score, f1_score

from .embedding_incremental import DEFAULT_C_GRID
from .interviewer_signal_explanation import (
    BASE_SEED,
    D_P_FEATURES,
    LABEL_C_GRID,
    RIDGE_ALPHAS,
    ALL_BLOCK_SUBSETS,
    _classification_metrics,
    _make_derangement,
    _paired_probability_delta_metrics,
    _subset_design,
    explanation_metrics,
    fit_explanation_fold,
    fit_label_outer,
    load_explanation_inputs,
    run_source_score_crossfit,
)
from .io import atomic_write_csv, atomic_write_json, sha256_file
from .splits import split_indices


BASE_FORMAL_RESULT_COMMIT = "ed66d68aa20e789fbd4829b44c49362b784a61e2"
C5_CANONICAL_RELATIVE = Path(
    "processed_research/latest_usable_data_official142_20260702/"
    "04_c5_evidence_sources/c5_quotes_only_v3_gpt55_reviewed_spans_official142.csv"
)
C5_CANONICAL_SHA256 = "794d7b43ebd95d57ac4a90a4217953a9bea27e98ff91b13c31b0a06f9aa48c7e"
LOCKED_REVISION_INPUT_HASHES = {
    "formal_explanation_subset_metrics": "a550ba89d14e0f470d87327002bc7e5cf2c3f8cb9c5f5e691de3fefe72c1b726",
    "formal_source_score_crossfit": "01ba43c969ec0e2cbbec8214b8dc90e6869cd1118339a237f4914a79e9d3dd80",
    "formal_run_manifest": "e665d0b003bd4ed22e1880901720d02ea7d58c198c3b3879bdb6c807f0533394",
    "formal_fake_domain_explanation_results": "682da41d7c873d3b3fef75e764961cc351e769edc9a9a0b264c451d552c50df0",
    "formal_fake_domain_residual_results": "6ac82cb7eb5af954d8cda70e082f2aca8edc8e93d30b3f0fa7648691990950f3",
    "c5_traceability_audit": "2a11e33d8996c98b6ba3f40ca406e67188ed77b8f09117d1bde4490e3649b186",
    "canonical_c5_spans": C5_CANONICAL_SHA256,
}
LOCKED_REVISION_INPUT_PATHS: dict[str, tuple[str, Path]] = {
    "formal_explanation_subset_metrics": (
        "output_root",
        Path("15_interviewer_signal_explanation/final/explanation_subset_metrics.csv"),
    ),
    "formal_source_score_crossfit": (
        "output_root",
        Path("15_interviewer_signal_explanation/final/source_score_crossfit.csv"),
    ),
    "formal_run_manifest": (
        "output_root",
        Path("15_interviewer_signal_explanation/final/run_manifest.json"),
    ),
    "formal_fake_domain_explanation_results": (
        "output_root",
        Path("15_interviewer_signal_explanation/final/fake_domain_explanation_results.csv"),
    ),
    "formal_fake_domain_residual_results": (
        "output_root",
        Path("15_interviewer_signal_explanation/final/fake_domain_residual_results.csv"),
    ),
    "c5_traceability_audit": (
        "output_root",
        Path("12_submission_audit/c5_traceability_audit_no_quotes.csv"),
    ),
    "canonical_c5_spans": ("project_root", C5_CANONICAL_RELATIVE),
}
C5_HUMAN_FIELDS = (
    "human_span_valid",
    "human_domain",
    "human_polarity",
    "human_false_positive",
    "human_missing_evidence",
    "rater_id",
    "notes",
)
C5_REVIEWER_FIELDS = (
    "review_case_id",
    "deidentified_quote",
    "local_context_reference",
    "human_span_valid",
    "human_domain",
    "human_polarity",
    "rater_id",
    "notes",
)
C5_MISSING_EVIDENCE_FIELDS = (
    "fulltext_review_case_id",
    "local_context_reference",
    "human_missing_evidence",
    "missing_evidence_quote",
    "rater_id",
    "notes",
)
C5_REVIEWER_PROHIBITED_FIELDS = {
    "participant_id",
    "label",
    "span_id",
    "source_order",
    "model",
    "model_domain",
    "model_polarity",
    "model_exact_quote",
    "model_match_status",
    "model_review_status",
    "model_original_quote",
    "source_split_reference",
    "source_sha256",
    "source_text_sha256",
    "quote_sha256",
    "canonical_input_sha256",
    "candidate_scope",
    "human_false_positive",
    "human_missing_evidence",
}
C5_MODEL_FIELDS = (
    "participant_id",
    "label",
    "span_id",
    "source_order",
    "model",
    "model_domain",
    "model_polarity",
    "model_exact_quote",
    "model_match_status",
    "model_review_status",
    "model_original_quote",
    "source_split_reference",
    "source_sha256",
    "source_text_sha256",
    "quote_sha256",
    "canonical_input_sha256",
    "candidate_scope",
)
R_ABSORPTION_MODELS = (
    "P+Q+D",
    "P+Q+D+residual_without_R",
    "P+Q+R+D",
    "P+Q+R+D+residual_full",
)
R_ABSORPTION_ENHANCEMENTS = {
    "P+Q+D+residual_without_R": ("P+Q+D", "residual_without_R"),
    "P+Q+R+D+residual_full": ("P+Q+R+D", "residual_full"),
}
SOURCE_PERFORMANCE_SOURCES = {
    "P": "participant_probability",
    "I": "interviewer_probability",
}
LOCKED_INCREMENT_SUBSETS = (
    "P",
    "Q",
    "R",
    "D",
    "P+R",
    "R+D",
    "P+R+D",
    "P+Q+R+D",
)
FIXED_XLSX_DATETIME = dt.datetime(2000, 1, 1, 0, 0, 0)


def _stable_json_hash(value: object) -> str:
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def _resolve_code_commit(repo_root: Path, explicit: str | None = None) -> str:
    if explicit:
        return str(explicit)
    try:
        return subprocess.check_output(
            ["git", "-C", str(repo_root), "rev-parse", "HEAD"],
            text=True,
            stderr=subprocess.DEVNULL,
        ).strip()
    except (OSError, subprocess.CalledProcessError):
        return "unknown"


def _clean(value: object) -> str:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return ""
    return str(value)


def find_c5_canonical_source(project_root: str | Path) -> Path:
    root = Path(project_root).resolve()
    direct = root / C5_CANONICAL_RELATIVE
    candidates = [direct]
    if not direct.exists():
        candidates.extend(
            sorted(
                root.joinpath("processed_research").rglob(
                    "c5_quotes_only_v3_gpt55_reviewed_spans_official142.csv"
                )
            )
        )
    for candidate in candidates:
        if candidate.exists() and sha256_file(candidate) == C5_CANONICAL_SHA256:
            return candidate
    observed = [str(candidate) for candidate in candidates if candidate.exists()]
    raise FileNotFoundError(
        "the locked official-142 C5 span source was not found with the expected SHA-256; "
        f"observed candidates={observed}"
    )


def validate_locked_revision_inputs(
    output_root: str | Path,
    project_root: str | Path,
) -> dict[str, str]:
    """Verify every frozen Stage 15 and C5 input before formal output exists."""

    roots = {
        "output_root": Path(output_root).resolve(),
        "project_root": Path(project_root).resolve(),
    }
    observed_hashes: dict[str, str] = {}
    for name, expected in LOCKED_REVISION_INPUT_HASHES.items():
        if name not in LOCKED_REVISION_INPUT_PATHS:
            raise ValueError(f"locked revision input path is not configured: {name}")
        root_name, relative_path = LOCKED_REVISION_INPUT_PATHS[name]
        if root_name not in roots:
            raise ValueError(f"locked revision input has unsupported root {root_name!r}: {name}")
        path = roots[root_name] / relative_path
        if not path.exists():
            raise FileNotFoundError(f"locked revision input not found for {name}: {path}")
        observed = sha256_file(path)
        if observed != expected:
            raise ValueError(
                "locked revision input hash mismatch for "
                f"{name}: expected={expected}, observed={observed}, path={path}"
            )
        observed_hashes[name] = observed
    return observed_hashes


def resolve_revision_project_root(
    output_root: str | Path,
    project_root: str | Path | None = None,
) -> Path:
    """Locate the repository-level canonical C5 source from a linked worktree."""

    if project_root is not None:
        candidates = [Path(project_root).resolve()]
    else:
        output = Path(output_root).resolve()
        candidates = [output.parent, *output.parent.parents]
    for candidate in candidates:
        if (candidate / C5_CANONICAL_RELATIVE).exists():
            return candidate
    raise FileNotFoundError(
        "could not locate the canonical C5 span source from the revision worktree; "
        f"checked={[str(candidate / C5_CANONICAL_RELATIVE) for candidate in candidates]}"
    )


def _tertile(values: pd.Series) -> pd.Series:
    numeric = pd.to_numeric(values, errors="raise").astype(float)
    ranks = numeric.rank(method="first")
    return pd.qcut(ranks, q=3, labels=False, duplicates="drop").astype(int) + 1


def select_c5_validation_participants(
    structural: pd.DataFrame,
    domains: pd.DataFrame,
    *,
    n: int = 24,
    evidence_counts: Mapping[object, object] | pd.Series | None = None,
) -> pd.DataFrame:
    """Select a deterministic label/length/D-P-breadth/evidence roster."""

    required = {"participant_id", "label", "length_only__participant_word_count"}
    missing = sorted(required.difference(structural.columns))
    if missing:
        raise ValueError(f"structural input is missing sampling fields: {missing}")
    missing_domains = sorted(set(D_P_FEATURES).difference(domains.columns))
    if missing_domains:
        raise ValueError(f"D-P input is missing sampling fields: {missing_domains}")
    if int(n) <= 0:
        raise ValueError("n must be positive")
    frame = structural[["participant_id", "label", "length_only__participant_word_count"]].copy()
    frame["participant_id"] = pd.to_numeric(frame["participant_id"], errors="raise").astype(int)
    frame["label"] = pd.to_numeric(frame["label"], errors="raise").astype(int)
    domain_frame = domains[["participant_id", *D_P_FEATURES]].copy()
    domain_frame["participant_id"] = pd.to_numeric(domain_frame["participant_id"], errors="raise").astype(int)
    frame = frame.merge(domain_frame, on="participant_id", how="inner", validate="one_to_one")
    frame["d_p_domain_breadth"] = (frame[list(D_P_FEATURES)] > 0).sum(axis=1).astype(int)
    if evidence_counts is None:
        evidence_lookup: Mapping[object, object] = {}
    elif isinstance(evidence_counts, pd.Series):
        evidence_lookup = evidence_counts.to_dict()
    else:
        evidence_lookup = evidence_counts
    frame["c5_evidence_count"] = frame["participant_id"].map(evidence_lookup).fillna(0).astype(int)
    frame["length_tertile"] = _tertile(frame["length_only__participant_word_count"])
    frame["domain_breadth_tertile"] = _tertile(frame["d_p_domain_breadth"])
    frame["evidence_count_tertile"] = _tertile(frame["c5_evidence_count"])
    frame["sampling_stratum"] = (
        "label=" + frame["label"].astype(str)
        + ";length_tertile=" + frame["length_tertile"].astype(str)
        + ";d_p_breadth_tertile=" + frame["domain_breadth_tertile"].astype(str)
        + ";evidence_tertile=" + frame["evidence_count_tertile"].astype(str)
    )
    frame = frame.sort_values(["sampling_stratum", "participant_id"], kind="stable").reset_index(drop=True)
    groups = [group.copy() for _, group in frame.groupby("sampling_stratum", sort=True)]
    selected: list[pd.Series] = []
    pointer = 0
    while len(selected) < min(int(n), len(frame)):
        progressed = False
        for group in groups:
            if pointer < len(group) and len(selected) < int(n):
                selected.append(group.iloc[pointer])
                progressed = True
        if not progressed:
            pointer += 1
    result = pd.DataFrame(selected).reset_index(drop=True)
    result.insert(0, "selection_order", np.arange(1, len(result) + 1, dtype=int))
    result["selection_rule"] = (
        "deterministic round-robin over label, participant-word-count tertile, "
        "D-P domain-breadth tertile, and C5 evidence-count tertile; "
        "no model correctness or source score"
    )
    return result[
        [
            "selection_order",
            "participant_id",
            "label",
            "length_tertile",
            "d_p_domain_breadth",
            "domain_breadth_tertile",
            "c5_evidence_count",
            "evidence_count_tertile",
            "sampling_stratum",
            "selection_rule",
        ]
    ]


def build_c5_human_validation_dataset(
    spans: pd.DataFrame,
    roster: pd.DataFrame,
    *,
    source_sha256: str,
    traceability: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Build the owner-only candidate sheet with model evidence and blank human fields."""

    required = {"participant_id", "exact_quote", "domain", "polarity"}
    missing = sorted(required.difference(spans.columns))
    if missing:
        raise ValueError(f"C5 span input is missing required model fields: {missing}")
    if "participant_id" not in roster.columns:
        raise ValueError("C5 roster is missing participant_id")
    sample_ids = pd.to_numeric(roster["participant_id"], errors="raise").astype(int).tolist()
    label_lookup = {}
    if "label" in roster.columns:
        label_lookup = dict(
            zip(
                pd.to_numeric(roster["participant_id"], errors="raise").astype(int),
                pd.to_numeric(roster["label"], errors="raise").astype(int),
            )
        )
    span_frame = spans.copy()
    span_frame["participant_id"] = pd.to_numeric(span_frame["participant_id"], errors="raise").astype(int)
    if "source_order" in span_frame.columns:
        span_frame["source_order"] = pd.to_numeric(span_frame["source_order"], errors="coerce")
    else:
        span_frame["source_order"] = np.nan
    if traceability is not None and not traceability.empty:
        trace = traceability.copy()
        trace["participant_id"] = pd.to_numeric(trace["participant_id"], errors="raise").astype(int)
        trace["source_order"] = pd.to_numeric(trace["source_order"], errors="raise")
        trace_columns = [
            column
            for column in ("source_text_sha256", "quote_sha256", "normalized_char_start", "normalized_char_end")
            if column in trace.columns
        ]
        if trace_columns:
            span_frame = span_frame.merge(
                trace[["participant_id", "source_order", *trace_columns]],
                on=["participant_id", "source_order"],
                how="left",
                validate="many_to_one",
                suffixes=("", "_trace"),
            )
    span_frame = span_frame[span_frame["participant_id"].isin(sample_ids)].copy()
    span_frame["_stable_order"] = np.arange(len(span_frame))
    span_frame = span_frame.sort_values(
        ["participant_id", "source_order", "_stable_order"], kind="stable"
    ).reset_index(drop=True)
    rows: list[dict[str, object]] = []
    for participant_id in sample_ids:
        participant_rows = span_frame.loc[span_frame["participant_id"].eq(int(participant_id))]
        if participant_rows.empty:
            participant_rows = pd.DataFrame([{"participant_id": int(participant_id), "_no_model_span": True}])
        for local_order, (_, row) in enumerate(participant_rows.iterrows(), start=1):
            no_model_span = bool(row.get("_no_model_span", False))
            source_order = "" if no_model_span else _clean(row.get("source_order"))
            rows.append(
                {
                    "participant_id": int(participant_id),
                    "label": label_lookup.get(int(participant_id), ""),
                    "span_id": f"{int(participant_id)}:{'NO_MODEL_EVIDENCE' if no_model_span else f'{local_order:04d}'}",
                    "source_order": source_order,
                    "model": _clean(row.get("model")),
                    "model_domain": "" if no_model_span else _clean(row.get("domain")),
                    "model_polarity": "" if no_model_span else _clean(row.get("polarity")),
                    "model_exact_quote": "" if no_model_span else _clean(row.get("exact_quote")),
                    "model_match_status": "" if no_model_span else _clean(row.get("match_status")),
                    "model_review_status": "" if no_model_span else _clean(row.get("review_status")),
                    "model_original_quote": "" if no_model_span else _clean(row.get("original_model_quote")),
                    "source_split_reference": "" if no_model_span else _clean(row.get("source_split_in_original_table")),
                    "source_sha256": str(source_sha256),
                    "source_text_sha256": "" if no_model_span else _clean(row.get("source_text_sha256")),
                    "quote_sha256": "" if no_model_span else _clean(row.get("quote_sha256")),
                    "canonical_input_sha256": str(source_sha256),
                    "candidate_scope": "C5 model-evidence candidate for independent human validation; transcript not copied",
                    **{field: "" for field in C5_HUMAN_FIELDS},
                }
            )
    result = pd.DataFrame(rows, columns=[*C5_MODEL_FIELDS, *C5_HUMAN_FIELDS])
    forbidden = {"transcript", "speaker_text", "participant_text", "interviewer_text"}
    if forbidden.intersection({str(column).lower() for column in result.columns}):
        raise AssertionError("C5 candidate contains a transcript field")
    return result


def _opaque_c5_review_id(prefix: str, *parts: object) -> str:
    return f"{prefix}-{_stable_json_hash([_clean(part) for part in parts])[:16]}"


def _deidentify_c5_quote(quote: object) -> str:
    """Apply limited deterministic redaction before a quote enters a rater packet."""

    value = _clean(quote)
    value = re.sub(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b", "[EMAIL]", value)
    value = re.sub(r"\b(?:\+?\d{1,2}[ .-]?)?(?:\(?\d{3}\)?[ .-]?)\d{3}[ .-]?\d{4}\b", "[PHONE]", value)
    value = re.sub(r"(?i)\b(my name is|i am|i'm)\s+[A-Z][a-z]+(?:\s+[A-Z][a-z]+)?", r"\1 [NAME]", value)
    value = re.sub(r"\b\d{3,}\b", "[NUMBER]", value)
    return re.sub(r"\s+", " ", value).strip()


def build_c5_owner_crosswalk(
    spans: pd.DataFrame,
    roster: pd.DataFrame,
    *,
    source_sha256: str,
    traceability: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Build the restricted owner-only C5 linkage and model-reference table."""

    owner = build_c5_human_validation_dataset(
        spans,
        roster,
        source_sha256=source_sha256,
        traceability=traceability,
    ).copy()
    owner.insert(
        0,
        "review_case_id",
        [
            _opaque_c5_review_id(
                "c5-review",
                source_sha256,
                row.participant_id,
                row.span_id,
                row.quote_sha256,
            )
            for row in owner.itertuples(index=False)
        ],
    )
    owner.insert(
        1,
        "fulltext_review_case_id",
        [
            _opaque_c5_review_id("c5-fulltext", source_sha256, participant_id)
            for participant_id in owner["participant_id"]
        ],
    )
    if owner["review_case_id"].duplicated().any():
        raise AssertionError("C5 owner crosswalk generated duplicate review_case_id values")
    return owner


def build_c5_blinded_reviewer_packet(owner_crosswalk: pd.DataFrame) -> pd.DataFrame:
    """Build a blinded candidate-span packet without participant or model labels."""

    required = {"review_case_id", "fulltext_review_case_id", "model_exact_quote"}
    missing = sorted(required.difference(owner_crosswalk.columns))
    if missing:
        raise ValueError(f"C5 owner crosswalk is missing blinded-packet fields: {missing}")
    candidates = owner_crosswalk.loc[
        owner_crosswalk["model_exact_quote"].astype(str).str.strip().ne("")
    ].copy()
    packet = pd.DataFrame(
        {
            "review_case_id": candidates["review_case_id"].astype(str),
            "deidentified_quote": candidates["model_exact_quote"].map(_deidentify_c5_quote),
            "local_context_reference": "restricted-fulltext/" + candidates["fulltext_review_case_id"].astype(str),
            "human_span_valid": "",
            "human_domain": "",
            "human_polarity": "",
            "rater_id": "",
            "notes": "",
        }
    )
    if packet["review_case_id"].duplicated().any():
        raise AssertionError("C5 reviewer packet generated duplicate review_case_id values")
    if C5_REVIEWER_PROHIBITED_FIELDS.intersection(packet.columns):
        raise AssertionError("C5 reviewer packet contains a prohibited owner/model field")
    return packet.loc[:, list(C5_REVIEWER_FIELDS)].reset_index(drop=True)


def build_c5_missing_evidence_audit_form(owner_crosswalk: pd.DataFrame) -> pd.DataFrame:
    """Build a separate participant-level full-text missing-evidence form."""

    required = {"fulltext_review_case_id"}
    missing = sorted(required.difference(owner_crosswalk.columns))
    if missing:
        raise ValueError(f"C5 owner crosswalk is missing full-text audit fields: {missing}")
    cases = owner_crosswalk[["fulltext_review_case_id"]].drop_duplicates().sort_values(
        "fulltext_review_case_id", kind="stable"
    )
    form = pd.DataFrame(
        {
            "fulltext_review_case_id": cases["fulltext_review_case_id"].astype(str).to_numpy(),
            "local_context_reference": (
                "restricted-fulltext/" + cases["fulltext_review_case_id"].astype(str)
            ).to_numpy(),
            "human_missing_evidence": "",
            "missing_evidence_quote": "",
            "rater_id": "",
            "notes": "",
        }
    )
    return form.loc[:, list(C5_MISSING_EVIDENCE_FIELDS)]


def _parse_human_binary(series: pd.Series, *, field: str) -> np.ndarray:
    normalized = series.astype(str).str.strip().str.lower()
    mapping = {
        "1": 1,
        "0": 0,
        "true": 1,
        "false": 0,
        "yes": 1,
        "no": 0,
        "valid": 1,
        "invalid": 0,
        "present": 1,
        "absent": 0,
    }
    values = normalized.map(mapping)
    if values.isna().any():
        raise ValueError(f"C5 human field {field} contains blank or unsupported labels")
    return values.to_numpy(dtype=int)


def _prepare_completed_c5_rater(frame: pd.DataFrame, *, name: str) -> pd.DataFrame:
    required = {"review_case_id", "human_span_valid", "human_domain", "human_polarity"}
    missing = sorted(required.difference(frame.columns))
    if missing:
        raise ValueError(f"{name} is missing blinded C5 human-validation fields: {missing}")
    result = frame.copy()
    if result["review_case_id"].isna().any() or result["review_case_id"].astype(str).str.strip().eq("").any():
        raise ValueError(f"{name} has blank review_case_id values")
    if result["review_case_id"].duplicated().any():
        raise ValueError(f"{name} has duplicate review_case_id rows")
    try:
        result["_human_span_valid"] = _parse_human_binary(
            result["human_span_valid"], field="human_span_valid"
        )
    except ValueError as exc:
        raise ValueError(f"{name} has incomplete human labels; cannot compute kappa or F1") from exc
    for field in ("human_domain", "human_polarity"):
        result[field] = result[field].map(_clean).str.strip()
    missing_valid_details = result["_human_span_valid"].eq(1) & (
        result["human_domain"].eq("") | result["human_polarity"].eq("")
    )
    if missing_valid_details.any():
        raise ValueError(f"{name} has a valid span without domain/polarity labels")
    return result.sort_values("review_case_id", kind="stable").reset_index(drop=True)


def _prepare_c5_owner_crosswalk(owner_crosswalk: pd.DataFrame, review_case_ids: pd.Series) -> pd.DataFrame:
    required = {"review_case_id", "model_exact_quote", "model_domain", "model_polarity"}
    missing = sorted(required.difference(owner_crosswalk.columns))
    if missing:
        raise ValueError(f"owner_crosswalk is missing model-reference fields: {missing}")
    owner = owner_crosswalk.copy()
    if owner["review_case_id"].duplicated().any():
        raise ValueError("owner_crosswalk has duplicate review_case_id rows")
    owner = owner.set_index("review_case_id", drop=False)
    missing_cases = sorted(set(review_case_ids.astype(str)).difference(owner.index.astype(str)))
    if missing_cases:
        raise ValueError(f"owner_crosswalk is missing rater review_case_id values: {missing_cases[:5]}")
    return owner.loc[review_case_ids.astype(str)].reset_index(drop=True)


def _safe_cohen_kappa(left: Sequence[object], right: Sequence[object]) -> float:
    left_values = np.asarray(list(left))
    right_values = np.asarray(list(right))
    if len(left_values) < 2 or len(right_values) < 2:
        return float("nan")
    labels = set(left_values.tolist()).union(right_values.tolist())
    if len(labels) < 2:
        return float("nan")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", category=UndefinedMetricWarning)
        warnings.simplefilter("ignore", category=UserWarning)
        return float(cohen_kappa_score(left_values, right_values))


def _safe_f1(
    left: Sequence[object],
    right: Sequence[object],
    *,
    average: str | None = None,
) -> float:
    left_values = np.asarray(list(left))
    right_values = np.asarray(list(right))
    if len(left_values) == 0 or len(right_values) == 0:
        return float("nan")
    kwargs: dict[str, object] = {"zero_division": 0}
    if average is not None:
        kwargs["average"] = average
        kwargs["labels"] = sorted(set(left_values.tolist()).union(right_values.tolist()))
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", category=UndefinedMetricWarning)
        warnings.simplefilter("ignore", category=UserWarning)
        return float(f1_score(left_values, right_values, **kwargs))


def compute_c5_human_machine_agreement(
    rater_a: pd.DataFrame,
    rater_b: pd.DataFrame,
    *,
    owner_crosswalk: pd.DataFrame | None = None,
    consensus: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Compute blinded C5 agreement under the locked candidate-span contract."""

    left = _prepare_completed_c5_rater(rater_a, name="rater_a")
    right = _prepare_completed_c5_rater(rater_b, name="rater_b")
    if not left["review_case_id"].eq(right["review_case_id"]).all():
        raise ValueError("C5 rater files do not contain the same candidate spans")
    rows: list[dict[str, object]] = []

    def append_metric(comparison: str, metric: str, value: float, n: int, scope: str) -> None:
        rows.append(
            {
                "comparison": comparison,
                "metric": metric,
                "estimate": float(value),
                "n": int(n),
                "scope": scope,
            }
        )

    append_metric(
        "rater_a_vs_rater_b",
        "span_valid_cohen_kappa",
        _safe_cohen_kappa(left["_human_span_valid"], right["_human_span_valid"]),
        len(left),
        "all_candidate_spans_independent_human_rater_agreement",
    )
    jointly_valid = left["_human_span_valid"].eq(1) & right["_human_span_valid"].eq(1)
    for field, metric_name in (("human_domain", "domain"), ("human_polarity", "polarity")):
        if int(jointly_valid.sum()) == 0:
            continue
        append_metric(
            "rater_a_vs_rater_b",
            f"{metric_name}_cohen_kappa",
            _safe_cohen_kappa(
                left.loc[jointly_valid, field],
                right.loc[jointly_valid, field],
            ),
            int(jointly_valid.sum()),
            "candidate_spans_valid_by_both_human_raters",
        )

    if owner_crosswalk is None:
        if consensus is not None:
            raise ValueError("owner_crosswalk is required for model-versus-consensus agreement")
        return pd.DataFrame(rows)
    owner = _prepare_c5_owner_crosswalk(owner_crosswalk, left["review_case_id"])

    def append_model_metrics(comparison: str, human: pd.DataFrame, rater_scope: str) -> None:
        model_present = owner["model_exact_quote"].map(_clean).str.strip().ne("").to_numpy(dtype=int)
        human_valid = human["_human_span_valid"].to_numpy(dtype=int)
        append_metric(
            comparison,
            "span_valid_f1",
            _safe_f1(human_valid, model_present),
            len(human),
            "sampled_candidate_span_presence_not_full_transcript_recall",
        )
        append_metric(
            comparison,
            "span_valid_cohen_kappa",
            _safe_cohen_kappa(human_valid, model_present),
            len(human),
            "sampled_candidate_span_presence_not_full_transcript_recall",
        )
        valid = human["_human_span_valid"].eq(1)
        for field, metric_name in (("domain", "domain"), ("polarity", "polarity")):
            if int(valid.sum()) == 0:
                continue
            model_values = owner.loc[valid, f"model_{field}"].map(_clean).str.strip()
            human_values = human.loc[valid, f"human_{field}"]
            append_metric(
                comparison,
                f"{metric_name}_macro_f1",
                _safe_f1(human_values, model_values, average="macro"),
                int(valid.sum()),
                rater_scope,
            )
            append_metric(
                comparison,
                f"{metric_name}_cohen_kappa",
                _safe_cohen_kappa(human_values, model_values),
                int(valid.sum()),
                rater_scope,
            )

    append_model_metrics("model_vs_rater_a", left, "rater_a_valid_candidate_spans")
    append_model_metrics("model_vs_rater_b", right, "rater_b_valid_candidate_spans")
    if consensus is not None:
        consensus_frame = _prepare_completed_c5_rater(consensus, name="consensus")
        if not left["review_case_id"].eq(consensus_frame["review_case_id"]).all():
            raise ValueError("C5 consensus file does not contain the same candidate spans")
        append_model_metrics("model_vs_consensus", consensus_frame, "consensus_valid_candidate_spans")
    return pd.DataFrame(rows)


def build_feature_block_source_table(inputs: object) -> pd.DataFrame:
    """Describe where P/Q/R/D variables come from without causal language."""

    block_columns = getattr(inputs, "block_columns")
    rows: list[dict[str, object]] = []
    for block in ("P", "Q", "R", "D"):
        columns = list(block_columns.get(block, ()))
        for feature in columns:
            if block == "P":
                side = "participant-side source model"
                source_file = "analysis_v2/15_interviewer_signal_explanation/final/source_score_crossfit.csv"
                definition = "strict cross-fitted participant source logit"
            elif block == "Q":
                side = "mixed participant/interviewer/process quantity"
                source_file = "analysis_v2/12_submission_audit/structural_baseline_features.csv"
                definition = "frozen length and interaction distribution feature"
            elif block == "R":
                side = "interviewer-side protocol structure"
                source_file = "structural baseline plus frozen interviewer path coverage features"
                definition = "frozen protocol count, prompt coverage, or audited path/prompt feature"
            else:
                side = "participant-derived D-P symptom-domain count"
                source_file = "analysis_v2/04_c5_controls/domain_count/input.csv"
                definition = "one of the ten frozen D-P symptom-domain counts"
            rows.append(
                {
                    "block": block,
                    "feature": str(feature),
                    "source_side": side,
                    "source_file": source_file,
                    "operational_definition": definition,
                    "uses_phq_label": False,
                    "uses_source_score": block == "P",
                    "interpretation_scope": "descriptive statistical recoverability source mapping",
                }
            )
    return pd.DataFrame(rows)


def build_block_increment_contrasts(metrics: pd.DataFrame) -> pd.DataFrame:
    """Compute descriptive R-block contrasts from frozen 16-subset R2 values."""

    required_columns = {"representation", "repeat", "subset", "r2"}
    missing = sorted(required_columns.difference(metrics.columns))
    if missing:
        raise ValueError(f"subset metrics are missing required columns: {missing}")
    rows: list[dict[str, object]] = []
    for (representation, repeat), group in metrics.groupby(["representation", "repeat"], sort=True):
        values = group.groupby("subset", sort=False)["r2"].first().to_dict()
        missing_subsets = [subset for subset in LOCKED_INCREMENT_SUBSETS if subset not in values]
        if missing_subsets:
            raise ValueError(
                f"missing locked subset(s) for representation={representation}, repeat={repeat}: {missing_subsets}"
            )
        full = float(values["P+Q+R+D"])
        r_alone = float(values["R"])
        rows.append(
            {
                "representation": str(representation),
                "repeat": int(repeat),
                "P_alone_r2": float(values["P"]),
                "Q_alone_r2": float(values["Q"]),
                "R_alone_r2": r_alone,
                "D_alone_r2": float(values["D"]),
                "full_r2": full,
                "full_minus_R_r2": full - r_alone,
                "P+R_minus_R_r2": float(values["P+R"]) - r_alone,
                "R+D_minus_R_r2": float(values["R+D"]) - r_alone,
                "P+R+D_minus_R_r2": float(values["P+R+D"]) - r_alone,
                "P+Q+R+D_minus_R_r2": full - r_alone,
                "interpretation_scope": "descriptive_score_recoverability_increment_not_significance",
            }
        )
    return pd.DataFrame(rows)


def _source_bootstrap_indices(labels: np.ndarray, *, n_bootstrap: int, seed: int) -> np.ndarray:
    labels = np.asarray(labels, dtype=int)
    positive = np.flatnonzero(labels == 1)
    negative = np.flatnonzero(labels == 0)
    if len(positive) == 0 or len(negative) == 0:
        raise ValueError("source performance bootstrap requires both label classes")
    rng = np.random.default_rng(int(seed))
    return np.asarray(
        [
            np.concatenate(
                [
                    rng.choice(negative, size=len(negative), replace=True),
                    rng.choice(positive, size=len(positive), replace=True),
                ]
            )
            for _ in range(int(n_bootstrap))
        ],
        dtype=int,
    )


def _validate_source_score_frame(
    source_scores: pd.DataFrame,
    *,
    expected_n: int = 142,
    expected_repeats: int = 10,
) -> pd.DataFrame:
    required = {
        "participant_id",
        "label",
        "representation",
        "repeat",
        "participant_probability",
        "interviewer_probability",
    }
    missing = sorted(required.difference(source_scores.columns))
    if missing:
        raise ValueError(f"source-score input is missing required columns: {missing}")
    frame = source_scores.copy()
    frame["participant_id"] = pd.to_numeric(frame["participant_id"], errors="raise").astype(int)
    frame["label"] = pd.to_numeric(frame["label"], errors="raise").astype(int)
    frame["repeat"] = pd.to_numeric(frame["repeat"], errors="raise").astype(int)
    for column in ("participant_probability", "interviewer_probability"):
        frame[column] = pd.to_numeric(frame[column], errors="raise")
    if not np.isfinite(frame[["participant_probability", "interviewer_probability"]].to_numpy(dtype=float)).all():
        raise ValueError("source-score probabilities contain non-finite values")
    if frame.duplicated(["participant_id", "representation", "repeat"]).any():
        raise ValueError("source-score input contains duplicate participant/repeat rows")
    for (representation, repeat), group in frame.groupby(["representation", "repeat"], sort=True):
        if len(group) != int(expected_n) or group["participant_id"].nunique() != int(expected_n):
            raise ValueError(
                f"source-score group {representation}/{repeat} does not contain the complete cohort"
            )
        if set(group["label"].unique()) != {0, 1} or int((group["label"] == 1).sum()) != 43:
            raise ValueError(f"source-score group {representation}/{repeat} does not preserve 43/99 labels")
    expected_repeats_set = set(range(1, int(expected_repeats) + 1))
    observed_repeats = set(frame["repeat"].unique())
    if not expected_repeats_set.issubset(observed_repeats):
        raise ValueError(f"source-score input is missing repeats: {sorted(expected_repeats_set - observed_repeats)}")
    return frame.sort_values(["representation", "repeat", "participant_id"], kind="stable").reset_index(drop=True)


def source_model_performance_by_repeat(
    source_scores: pd.DataFrame | str | Path,
    *,
    n_bootstrap: int = 5000,
    seed: int = BASE_SEED,
    expected_n: int = 142,
    expected_repeats: int = 10,
) -> pd.DataFrame:
    """Summarize frozen outer-test P/I source performance without refitting."""

    frame = pd.read_csv(source_scores) if isinstance(source_scores, (str, Path)) else source_scores.copy()
    frame = _validate_source_score_frame(frame, expected_n=expected_n, expected_repeats=expected_repeats)
    metric_names = (
        "roc_auc",
        "pr_auc",
        "brier",
        "log_loss",
        "sensitivity_at_0_50",
        "specificity_at_0_50",
    )
    rows: list[dict[str, object]] = []
    for (representation, repeat), group in frame.groupby(["representation", "repeat"], sort=True):
        labels = group["label"].to_numpy(dtype=int)
        for source_name, probability_column in SOURCE_PERFORMANCE_SOURCES.items():
            probability = group[probability_column].to_numpy(dtype=float)
            point = _classification_metrics(labels, probability)
            ci_fields = {f"{metric}_ci_low": float("nan") for metric in metric_names}
            ci_fields.update({f"{metric}_ci_high": float("nan") for metric in metric_names})
            bootstrap_scope = "not_run_repeat_gt_1"
            if int(repeat) == 1:
                index_draws = _source_bootstrap_indices(
                    labels,
                    n_bootstrap=n_bootstrap,
                    seed=int(seed) + sum(ord(char) for char in f"{representation}:{source_name}"),
                )
                distributions = {metric: [] for metric in metric_names}
                for index in index_draws:
                    sampled = _classification_metrics(labels[index], probability[index])
                    for metric in metric_names:
                        distributions[metric].append(float(sampled[metric]))
                for metric in metric_names:
                    values = np.asarray(distributions[metric], dtype=float)
                    ci_fields[f"{metric}_ci_low"] = float(np.quantile(values, 0.025))
                    ci_fields[f"{metric}_ci_high"] = float(np.quantile(values, 0.975))
                bootstrap_scope = "repeat_1_participant_stratified_bootstrap"
            rows.append(
                {
                    "representation": str(representation),
                    "repeat": int(repeat),
                    "source": source_name,
                    "n_participants": int(len(group)),
                    "n_positive": int((labels == 1).sum()),
                    "n_negative": int((labels == 0).sum()),
                    **point,
                    **ci_fields,
                    "n_bootstrap": int(n_bootstrap) if int(repeat) == 1 else 0,
                    "bootstrap_scope": bootstrap_scope,
                    "interpretation_scope": "descriptive_source_model_performance_outer_test_only",
                }
            )
    return pd.DataFrame(rows)


def source_model_performance_summary(performance_by_repeat: pd.DataFrame) -> pd.DataFrame:
    """Aggregate repeats descriptively; repeats are not treated as observations for tests."""

    required = {"representation", "repeat", "source"}
    missing = sorted(required.difference(performance_by_repeat.columns))
    if missing:
        raise ValueError(f"source performance table is missing columns: {missing}")
    metric_names = (
        "roc_auc",
        "pr_auc",
        "brier",
        "log_loss",
        "sensitivity_at_0_50",
        "specificity_at_0_50",
    )
    rows: list[dict[str, object]] = []
    for (representation, source), group in performance_by_repeat.groupby(["representation", "source"], sort=True):
        row: dict[str, object] = {
            "representation": str(representation),
            "source": str(source),
            "repeat_n": int(group["repeat"].nunique()),
            "repeats": ",".join(str(value) for value in sorted(group["repeat"].unique())),
            "summary_scope": "descriptive_repeat_summary_no_repeat_level_significance_test",
        }
        for metric in metric_names:
            values = group[metric].to_numpy(dtype=float)
            row[f"{metric}_mean"] = float(np.mean(values))
            row[f"{metric}_sd"] = float(np.std(values, ddof=1)) if len(values) > 1 else 0.0
            row[f"{metric}_median"] = float(np.median(values))
            row[f"{metric}_min"] = float(np.min(values))
            row[f"{metric}_max"] = float(np.max(values))
        rows.append(row)
    return pd.DataFrame(rows)


def apply_stratified_fake_domain_permutation(
    *,
    participant_ids: Sequence[object],
    labels: Sequence[int],
    domains: np.ndarray,
    train_idx: Sequence[int],
    test_idx: Sequence[int],
    draw: int,
    seed: int,
) -> tuple[np.ndarray, pd.DataFrame]:
    """Derange complete D-P rows within partition and label strata only."""

    ids = np.asarray(participant_ids)
    label_array = np.asarray(labels, dtype=int)
    matrix = np.asarray(domains, dtype=float)
    train_idx = np.asarray(train_idx, dtype=int)
    test_idx = np.asarray(test_idx, dtype=int)
    if matrix.shape[0] != len(ids) or label_array.shape[0] != len(ids):
        raise ValueError("stratified Fake-D inputs are not cohort aligned")
    if set(train_idx).intersection(test_idx):
        raise ValueError("stratified Fake-D train/test indices overlap")
    result = matrix.copy()
    rows: list[dict[str, object]] = []
    for partition_index, (partition, indices) in enumerate((("train", train_idx), ("test", test_idx))):
        for label_value in (0, 1):
            stratum = indices[label_array[indices] == int(label_value)]
            if len(stratum) < 2:
                raise ValueError(
                    f"stratified Fake-D {partition}/label={label_value} requires at least two participants"
                )
            local_permutation = _make_derangement(
                len(stratum),
                seed=int(seed) + int(draw) * 10_000 + partition_index * 100 + int(label_value),
            )
            donors = stratum[local_permutation]
            result[stratum] = matrix[donors]
            for recipient_index, donor_index in zip(stratum, donors):
                rows.append(
                    {
                        "draw": int(draw),
                        "partition": partition,
                        "label": int(label_value),
                        "recipient_index": int(recipient_index),
                        "donor_index": int(donor_index),
                        "recipient_id": ids[recipient_index],
                        "donor_id": ids[donor_index],
                        "self_match": bool(recipient_index == donor_index),
                        "same_label": bool(label_array[recipient_index] == label_array[donor_index]),
                    }
                )
    ledger = pd.DataFrame(rows).sort_values(
        ["partition", "label", "recipient_index"], kind="stable"
    ).reset_index(drop=True)
    ledger["donor_reuse_count"] = ledger.groupby(
        ["partition", "label", "donor_id"]
    )["donor_id"].transform("size").astype(int)
    ledger["one_to_one"] = ledger.groupby(["partition", "label"])["donor_id"].transform("nunique").eq(
        ledger.groupby(["partition", "label"])["donor_id"].transform("size")
    )
    if ledger["self_match"].any() or (~ledger["same_label"]).any() or (~ledger["one_to_one"]).any():
        raise AssertionError("stratified Fake-D assignment violated its partition/label contract")
    return result, ledger


def _fake_d_result_metrics(
    explanation_rows: list[dict[str, object]],
    *,
    representation: str,
    repeat: int,
    draw: int,
    domain_type: str,
) -> dict[str, object]:
    frame = pd.DataFrame(explanation_rows).sort_values("participant_id", kind="stable")
    metrics = explanation_metrics(
        frame["interviewer_logit"].to_numpy(dtype=float),
        frame["predicted_interviewer_logit"].to_numpy(dtype=float),
        frame["outer_train_mean_null_prediction"].to_numpy(dtype=float),
    )
    return {
        "representation": str(representation),
        "repeat": int(repeat),
        "draw": int(draw),
        "domain_type": str(domain_type),
        **metrics,
        "interpretation_scope": "targeted_post_result_sensitivity_real_D_vs_label_stratified_fake_D",
    }


def _stratified_fake_summary(
    explanation_oof: pd.DataFrame,
    residual_oof: pd.DataFrame,
    *,
    non_stratified_explanation_results: pd.DataFrame | None,
    non_stratified_residual_results: pd.DataFrame | None,
    n_bootstrap: int,
    seed: int,
) -> pd.DataFrame:
    real_explanation = explanation_oof.loc[explanation_oof["domain_type"].eq("real")].sort_values("participant_id", kind="stable")
    fake_explanation = explanation_oof.loc[explanation_oof["domain_type"].eq("fake")]
    real_residual = residual_oof.loc[residual_oof["domain_type"].eq("real")].sort_values("participant_id", kind="stable")
    fake_residual = residual_oof.loc[residual_oof["domain_type"].eq("fake")]
    if real_explanation.empty or fake_explanation.empty or real_residual.empty or fake_residual.empty:
        raise ValueError("stratified Fake-D summary requires real and fake OOF rows")
    participant_ids = real_explanation["participant_id"].to_numpy()
    labels = real_explanation["label"].to_numpy(dtype=int)
    draws = sorted(int(value) for value in fake_explanation["draw"].unique())
    fake_explanation_by_draw = {
        draw: fake_explanation.loc[fake_explanation["draw"].eq(draw)].sort_values("participant_id", kind="stable").reset_index(drop=True)
        for draw in draws
    }
    fake_residual_by_draw = {
        draw: fake_residual.loc[fake_residual["draw"].eq(draw)].sort_values("participant_id", kind="stable").reset_index(drop=True)
        for draw in draws
    }
    for frame in [*fake_explanation_by_draw.values(), *fake_residual_by_draw.values()]:
        if not np.array_equal(frame["participant_id"].to_numpy(), participant_ids):
            raise ValueError("stratified Fake-D OOF participant alignment failed")
    real_r2 = float(
        explanation_metrics(
            real_explanation["interviewer_logit"],
            real_explanation["predicted_interviewer_logit"],
            real_explanation["outer_train_mean_null_prediction"],
        )["r2"]
    )
    fake_r2 = np.asarray(
        [
            explanation_metrics(
                frame["interviewer_logit"],
                frame["predicted_interviewer_logit"],
                frame["outer_train_mean_null_prediction"],
            )["r2"]
            for frame in fake_explanation_by_draw.values()
        ],
        dtype=float,
    )
    real_residual_delta = _paired_probability_delta_metrics(
        real_residual["label"].to_numpy(dtype=int),
        real_residual["base_probability"].to_numpy(dtype=float),
        real_residual["enhanced_probability"].to_numpy(dtype=float),
    )["delta_auc"]
    fake_residual_delta = np.asarray(
        [
            _paired_probability_delta_metrics(
                frame["label"].to_numpy(dtype=int),
                frame["base_probability"].to_numpy(dtype=float),
                frame["enhanced_probability"].to_numpy(dtype=float),
            )["delta_auc"]
            for frame in fake_residual_by_draw.values()
        ],
        dtype=float,
    )
    rng = np.random.default_rng(int(seed))
    positive = np.flatnonzero(labels == 1)
    negative = np.flatnonzero(labels == 0)
    if len(positive) == 0 or len(negative) == 0:
        raise ValueError("stratified Fake-D summary requires both label classes")
    r2_distribution: list[float] = []
    residual_distribution: list[float] = []
    for _ in range(int(n_bootstrap)):
        index = np.concatenate(
            [
                rng.choice(negative, size=len(negative), replace=True),
                rng.choice(positive, size=len(positive), replace=True),
            ]
        )
        draw = draws[int(rng.integers(0, len(draws)))]
        fake_explanation_frame = fake_explanation_by_draw[draw]
        fake_residual_frame = fake_residual_by_draw[draw]
        real_r2_sample = explanation_metrics(
            real_explanation["interviewer_logit"].to_numpy()[index],
            real_explanation["predicted_interviewer_logit"].to_numpy()[index],
            real_explanation["outer_train_mean_null_prediction"].to_numpy()[index],
        )["r2"]
        fake_r2_sample = explanation_metrics(
            fake_explanation_frame["interviewer_logit"].to_numpy()[index],
            fake_explanation_frame["predicted_interviewer_logit"].to_numpy()[index],
            fake_explanation_frame["outer_train_mean_null_prediction"].to_numpy()[index],
        )["r2"]
        r2_distribution.append(float(real_r2_sample - fake_r2_sample))
        real_delta = _paired_probability_delta_metrics(
            labels[index],
            real_residual["base_probability"].to_numpy()[index],
            real_residual["enhanced_probability"].to_numpy()[index],
        )["delta_auc"]
        fake_delta = _paired_probability_delta_metrics(
            labels[index],
            fake_residual_frame["base_probability"].to_numpy()[index],
            fake_residual_frame["enhanced_probability"].to_numpy()[index],
        )["delta_auc"]
        residual_distribution.append(float(real_delta - fake_delta))
    r2_distribution = np.asarray(r2_distribution, dtype=float)
    residual_distribution = np.asarray(residual_distribution, dtype=float)
    observed_r2 = real_r2 - float(fake_r2.mean())
    observed_residual = float(real_residual_delta - fake_residual_delta.mean())
    nonstrat_r2 = float("nan")
    nonstrat_delta = float("nan")
    if non_stratified_explanation_results is not None and not non_stratified_explanation_results.empty:
        values = non_stratified_explanation_results.loc[
            non_stratified_explanation_results["domain_type"].eq("fake"), "r2"
        ]
        if not values.empty:
            nonstrat_r2 = float(values.mean())
    if non_stratified_residual_results is not None and not non_stratified_residual_results.empty:
        values = non_stratified_residual_results.loc[
            non_stratified_residual_results["domain_type"].eq("fake"), "delta_auc"
        ]
        if not values.empty:
            nonstrat_delta = float(values.mean())
    return pd.DataFrame(
        [
            {
                "representation": str(real_explanation["representation"].iloc[0]),
                "repeat": int(real_explanation["repeat"].iloc[0]),
                "real_D_r2": real_r2,
                "non_stratified_fake_D_r2_mean": nonstrat_r2,
                "stratified_fake_D_r2_mean": float(fake_r2.mean()),
                "stratified_fake_D_r2_sd": float(np.std(fake_r2, ddof=1)),
                "stratified_fake_D_r2_q025": float(np.quantile(fake_r2, 0.025)),
                "stratified_fake_D_r2_q975": float(np.quantile(fake_r2, 0.975)),
                "real_minus_stratified_fake_D_r2": observed_r2,
                "real_minus_stratified_fake_D_r2_ci_low": float(np.quantile(r2_distribution, 0.025)),
                "real_minus_stratified_fake_D_r2_ci_high": float(np.quantile(r2_distribution, 0.975)),
                "real_D_residual_delta_auc": float(real_residual_delta),
                "non_stratified_fake_D_residual_delta_auc_mean": nonstrat_delta,
                "stratified_fake_D_residual_delta_auc_mean": float(fake_residual_delta.mean()),
                "stratified_fake_D_residual_delta_auc_sd": float(np.std(fake_residual_delta, ddof=1)),
                "stratified_fake_D_residual_delta_auc_q025": float(np.quantile(fake_residual_delta, 0.025)),
                "stratified_fake_D_residual_delta_auc_q975": float(np.quantile(fake_residual_delta, 0.975)),
                "real_minus_stratified_fake_D_residual_delta_auc": observed_residual,
                "real_minus_stratified_fake_D_residual_delta_auc_ci_low": float(np.quantile(residual_distribution, 0.025)),
                "real_minus_stratified_fake_D_residual_delta_auc_ci_high": float(np.quantile(residual_distribution, 0.975)),
                "n_draws": int(len(draws)),
                "n_bootstrap": int(n_bootstrap),
                "inference": "effect estimates and participant-and-draw bootstrap intervals only",
                "interpretation_scope": "targeted_negative_control_sensitivity",
            }
        ]
    )


def targeted_fake_d_manifest_contract() -> dict[str, object]:
    """Return the locked non-inferential contract for label-stratified Fake-D."""

    return {
        "draws": 100,
        "partition_scope": "outer_train_and_outer_test_separately",
        "label_scope": "within_label_0_and_within_label_1",
        "rowwise_derangement": True,
        "cross_partition": False,
        "status": "targeted_post_result_sensitivity",
        "targeted_negative_control_sensitivity": True,
        "modifies_frozen_primary_FDR_family": False,
        "new_targeted_sensitivity_FDR_family": False,
        "inference": "effect estimates and participant-and-draw bootstrap intervals only",
    }


def run_stratified_fake_domain_controls(
    inputs: object,
    source_cache: Mapping[tuple[str, int, int], object],
    membership: pd.DataFrame,
    *,
    representation: str = "tfidf",
    repeat: int = 1,
    n_draws: int = 100,
    inner_folds: int = 3,
    alpha_grid: Sequence[float] = RIDGE_ALPHAS,
    label_c_grid: Sequence[float] = LABEL_C_GRID,
    base_seed: int = BASE_SEED,
    n_bootstrap: int = 5000,
    non_stratified_explanation_results: pd.DataFrame | None = None,
    non_stratified_residual_results: pd.DataFrame | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Run TF-IDF repeat-1 label-stratified Fake-D as a targeted sensitivity."""

    if int(n_draws) <= 0:
        raise ValueError("stratified Fake-D requires a positive draw count")
    domains = inputs.domains[list(D_P_FEATURES)].to_numpy(dtype=float)
    labels = np.asarray(inputs.labels, dtype=int)
    assignment_rows: list[dict[str, object]] = []
    explanation_oof_rows: list[dict[str, object]] = []
    residual_oof_rows: list[dict[str, object]] = []
    for fold in range(1, 6):
        key = (str(representation), int(repeat), int(fold))
        if key not in source_cache:
            raise ValueError(f"stratified Fake-D source cache is missing {key}")
        source = source_cache[key]
        train_idx = np.asarray(source.train_idx, dtype=int)
        test_idx = np.asarray(source.test_idx, dtype=int)
        target_train = np.asarray(source.interviewer_train_logit, dtype=float)
        target_test = np.asarray(source.interviewer_test_logit, dtype=float)
        fold_seed = int(base_seed) + int(repeat) * 100_000 + int(fold) * 1_000
        real_blocks = inputs.block_matrices
        real_train = _subset_design(
            "P+Q+R+D",
            participant_logits=source.participant_train_logit,
            block_matrices=real_blocks,
            indices=train_idx,
        )
        real_test = _subset_design(
            "P+Q+R+D",
            participant_logits=source.participant_test_logit,
            block_matrices=real_blocks,
            indices=test_idx,
        )
        real_explanation = fit_explanation_fold(
            real_train,
            target_train,
            real_test,
            stratify_labels=labels[train_idx],
            inner_folds=inner_folds,
            alpha_grid=alpha_grid,
            seed=fold_seed + 6_001,
        )
        real_residual_train = target_train - real_explanation.train_prediction
        real_residual_test = target_test - real_explanation.test_prediction
        real_base, _, _ = fit_label_outer(
            real_train,
            labels[train_idx],
            real_test,
            inner_folds=inner_folds,
            c_grid=label_c_grid,
            seed=fold_seed + 6_101,
        )
        real_enhanced, _, _ = fit_label_outer(
            np.column_stack([real_train, real_residual_train]),
            labels[train_idx],
            np.column_stack([real_test, real_residual_test]),
            inner_folds=inner_folds,
            c_grid=label_c_grid,
            seed=fold_seed + 6_201,
        )
        for local_index, global_index in enumerate(test_idx):
            explanation_oof_rows.append(
                {
                    "representation": str(representation),
                    "repeat": int(repeat),
                    "outer_fold": int(fold),
                    "draw": 0,
                    "domain_type": "real",
                    "participant_id": int(inputs.participant_ids[global_index]),
                    "label": int(labels[global_index]),
                    "interviewer_logit": float(target_test[local_index]),
                    "predicted_interviewer_logit": float(real_explanation.test_prediction[local_index]),
                    "outer_train_mean_null_prediction": float(np.mean(target_train)),
                }
            )
            residual_oof_rows.append(
                {
                    "representation": str(representation),
                    "repeat": int(repeat),
                    "outer_fold": int(fold),
                    "draw": 0,
                    "domain_type": "real",
                    "participant_id": int(inputs.participant_ids[global_index]),
                    "label": int(labels[global_index]),
                    "base_probability": float(real_base[local_index]),
                    "enhanced_probability": float(real_enhanced[local_index]),
                }
            )
        for draw in range(1, int(n_draws) + 1):
            fake_domains, ledger = apply_stratified_fake_domain_permutation(
                participant_ids=inputs.participant_ids,
                labels=labels,
                domains=domains,
                train_idx=train_idx,
                test_idx=test_idx,
                draw=draw,
                seed=fold_seed + 7_001,
            )
            ledger.insert(0, "representation", str(representation))
            ledger.insert(1, "repeat", int(repeat))
            ledger.insert(2, "outer_fold", int(fold))
            assignment_rows.extend(ledger.to_dict("records"))
            fake_blocks = dict(real_blocks)
            if all(column in D_P_FEATURES for column in inputs.block_columns["D"]):
                d_indices = [D_P_FEATURES.index(column) for column in inputs.block_columns["D"]]
            else:
                d_indices = list(range(inputs.block_matrices["D"].shape[1]))
            fake_blocks["D"] = fake_domains[:, d_indices]
            fake_train = _subset_design(
                "P+Q+R+D",
                participant_logits=source.participant_train_logit,
                block_matrices=fake_blocks,
                indices=train_idx,
            )
            fake_test = _subset_design(
                "P+Q+R+D",
                participant_logits=source.participant_test_logit,
                block_matrices=fake_blocks,
                indices=test_idx,
            )
            fake_explanation = fit_explanation_fold(
                fake_train,
                target_train,
                fake_test,
                stratify_labels=labels[train_idx],
                inner_folds=inner_folds,
                alpha_grid=alpha_grid,
                seed=fold_seed + 8_000 + draw,
            )
            fake_residual_train = target_train - fake_explanation.train_prediction
            fake_residual_test = target_test - fake_explanation.test_prediction
            fake_base, _, _ = fit_label_outer(
                fake_train,
                labels[train_idx],
                fake_test,
                inner_folds=inner_folds,
                c_grid=label_c_grid,
                seed=fold_seed + 9_000 + draw,
            )
            fake_enhanced, _, _ = fit_label_outer(
                np.column_stack([fake_train, fake_residual_train]),
                labels[train_idx],
                np.column_stack([fake_test, fake_residual_test]),
                inner_folds=inner_folds,
                c_grid=label_c_grid,
                seed=fold_seed + 10_000 + draw,
            )
            for local_index, global_index in enumerate(test_idx):
                explanation_oof_rows.append(
                    {
                        "representation": str(representation),
                        "repeat": int(repeat),
                        "outer_fold": int(fold),
                        "draw": int(draw),
                        "domain_type": "fake",
                        "participant_id": int(inputs.participant_ids[global_index]),
                        "label": int(labels[global_index]),
                        "interviewer_logit": float(target_test[local_index]),
                        "predicted_interviewer_logit": float(fake_explanation.test_prediction[local_index]),
                        "outer_train_mean_null_prediction": float(np.mean(target_train)),
                    }
                )
                residual_oof_rows.append(
                    {
                        "representation": str(representation),
                        "repeat": int(repeat),
                        "outer_fold": int(fold),
                        "draw": int(draw),
                        "domain_type": "fake",
                        "participant_id": int(inputs.participant_ids[global_index]),
                        "label": int(labels[global_index]),
                        "base_probability": float(fake_base[local_index]),
                        "enhanced_probability": float(fake_enhanced[local_index]),
                    }
                )
    explanation_oof = pd.DataFrame(explanation_oof_rows).sort_values(
        ["domain_type", "draw", "outer_fold", "participant_id"], kind="stable"
    ).reset_index(drop=True)
    residual_oof = pd.DataFrame(residual_oof_rows).sort_values(
        ["domain_type", "draw", "outer_fold", "participant_id"], kind="stable"
    ).reset_index(drop=True)
    explanation_results = pd.DataFrame(
        [
            _fake_d_result_metrics(
                group.to_dict("records"),
                representation=representation,
                repeat=repeat,
                draw=int(draw),
                domain_type=str(domain_type),
            )
            for (domain_type, draw), group in explanation_oof.groupby(["domain_type", "draw"], sort=True)
        ]
    )
    residual_results_rows: list[dict[str, object]] = []
    for (domain_type, draw), group in residual_oof.groupby(["domain_type", "draw"], sort=True):
        metric = _paired_probability_delta_metrics(
            group["label"].to_numpy(dtype=int),
            group["base_probability"].to_numpy(dtype=float),
            group["enhanced_probability"].to_numpy(dtype=float),
        )
        residual_results_rows.append(
            {
                "representation": str(representation),
                "repeat": int(repeat),
                "draw": int(draw),
                "domain_type": str(domain_type),
                **metric,
                "interpretation_scope": "targeted_post_result_sensitivity_real_D_vs_label_stratified_fake_D",
            }
        )
    residual_results = pd.DataFrame(residual_results_rows)
    summary = _stratified_fake_summary(
        explanation_oof,
        residual_oof,
        non_stratified_explanation_results=non_stratified_explanation_results,
        non_stratified_residual_results=non_stratified_residual_results,
        n_bootstrap=n_bootstrap,
        seed=base_seed + 70_000,
    )
    return assignment_rows_to_frame(assignment_rows), explanation_results, residual_results, summary, explanation_oof, residual_oof


def assignment_rows_to_frame(rows: list[dict[str, object]]) -> pd.DataFrame:
    """Return a stable empty/non-empty assignment table for output contracts."""

    columns = [
        "representation", "repeat", "outer_fold", "draw", "partition", "label",
        "recipient_index", "donor_index", "recipient_id", "donor_id", "self_match",
        "same_label", "donor_reuse_count", "one_to_one",
    ]
    return pd.DataFrame(rows, columns=columns).sort_values(
        ["representation", "repeat", "outer_fold", "draw", "partition", "label", "recipient_index"],
        kind="stable",
    ).reset_index(drop=True)


def _classification_metric_row(
    labels: np.ndarray,
    probability: np.ndarray,
    *,
    representation: str,
    repeat: int,
    model_name: str,
    delta_from: str = "",
) -> dict[str, object]:
    return {
        "representation": str(representation),
        "repeat": int(repeat),
        "model_name": model_name,
        "metric_scope": "label_prediction",
        "delta_from_model": delta_from,
        **_classification_metrics(labels, probability),
        "interpretation_scope": "exploratory_R_absorption_sensitivity_outside_primary_FDR",
    }


def run_r_absorption_sensitivity(
    inputs: object,
    membership: pd.DataFrame,
    output_root: str | Path,
    *,
    representation: str = "tfidf",
    repeat: int = 1,
    source_cache: Mapping[tuple[str, int, int], object] | None = None,
    inner_folds: int = 3,
    alpha_grid: Sequence[float] = RIDGE_ALPHAS,
    label_c_grid: Sequence[float] = LABEL_C_GRID,
    base_seed: int = BASE_SEED,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Run strict outer-fold R absorption using one cached source fit per fold."""

    if source_cache is None:
        _, source_cache, source_tuning = run_source_score_crossfit(
            inputs,
            membership,
            output_root,
            representations=(representation,),
            repeats=(int(repeat),),
            inner_folds=inner_folds,
            c_grid=DEFAULT_C_GRID,
            base_seed=base_seed,
        )
        tuning_rows: list[dict[str, object]] = list(source_tuning.to_dict("records"))
    else:
        tuning_rows = []
    prediction_rows: list[dict[str, object]] = []
    tuning_rows.extend([])
    for fold in range(1, 6):
        key = (str(representation), int(repeat), int(fold))
        if key not in source_cache:
            raise ValueError(f"R absorption source cache is missing {key}")
        source = source_cache[key]
        train_idx = np.asarray(source.train_idx, dtype=int)
        test_idx = np.asarray(source.test_idx, dtype=int)
        target_train = np.asarray(source.interviewer_train_logit, dtype=float)
        target_test = np.asarray(source.interviewer_test_logit, dtype=float)
        pqd_train = _subset_design(
            "P+Q+D",
            participant_logits=source.participant_train_logit,
            block_matrices=inputs.block_matrices,
            indices=train_idx,
        )
        pqd_test = _subset_design(
            "P+Q+D",
            participant_logits=source.participant_test_logit,
            block_matrices=inputs.block_matrices,
            indices=test_idx,
        )
        full_train = _subset_design(
            "P+Q+R+D",
            participant_logits=source.participant_train_logit,
            block_matrices=inputs.block_matrices,
            indices=train_idx,
        )
        full_test = _subset_design(
            "P+Q+R+D",
            participant_logits=source.participant_test_logit,
            block_matrices=inputs.block_matrices,
            indices=test_idx,
        )
        labels_train = np.asarray(inputs.labels, dtype=int)[train_idx]
        labels_test = np.asarray(inputs.labels, dtype=int)[test_idx]
        fold_seed = int(base_seed) + int(repeat) * 100_000 + int(fold) * 1_000
        pqd_explanation = fit_explanation_fold(
            pqd_train,
            target_train,
            pqd_test,
            stratify_labels=labels_train,
            inner_folds=inner_folds,
            alpha_grid=alpha_grid,
            seed=fold_seed + 701,
        )
        full_explanation = fit_explanation_fold(
            full_train,
            target_train,
            full_test,
            stratify_labels=labels_train,
            inner_folds=inner_folds,
            alpha_grid=alpha_grid,
            seed=fold_seed + 702,
        )
        residual_without_r_train = target_train - pqd_explanation.train_prediction
        residual_without_r_test = target_test - pqd_explanation.test_prediction
        residual_full_train = target_train - full_explanation.train_prediction
        residual_full_test = target_test - full_explanation.test_prediction
        label_models = (
            ("P+Q+D", pqd_train, pqd_test, None, None),
            (
                "P+Q+D+residual_without_R",
                np.column_stack([pqd_train, residual_without_r_train]),
                np.column_stack([pqd_test, residual_without_r_test]),
                "residual_without_R",
                residual_without_r_test,
            ),
            ("P+Q+R+D", full_train, full_test, None, None),
            (
                "P+Q+R+D+residual_full",
                np.column_stack([full_train, residual_full_train]),
                np.column_stack([full_test, residual_full_test]),
                "residual_full",
                residual_full_test,
            ),
        )
        probabilities: dict[str, np.ndarray] = {}
        selected_c: dict[str, float] = {}
        for model_index, (model_name, x_train, x_test, _, _) in enumerate(label_models):
            probability, selected, scores = fit_label_outer(
                x_train,
                labels_train,
                x_test,
                inner_folds=inner_folds,
                c_grid=label_c_grid,
                seed=fold_seed + 800 + model_index,
            )
            probabilities[model_name] = probability
            selected_c[model_name] = float(selected)
            for c_value, auc in scores.items():
                tuning_rows.append(
                    {
                        "representation": str(representation),
                        "repeat": int(repeat),
                        "outer_fold": int(fold),
                        "model_name": model_name,
                        "parameter": "C",
                        "value": float(c_value),
                        "inner_score_roc_auc": float(auc),
                        "selected": bool(float(c_value) == float(selected)),
                    }
                )
        for local_index, global_index in enumerate(test_idx):
            prediction_rows.append(
                {
                    "participant_id": int(inputs.participant_ids[global_index]),
                    "label": int(inputs.labels[global_index]),
                    "representation": str(representation),
                    "repeat": int(repeat),
                    "outer_fold": int(fold),
                    "interviewer_logit": float(target_test[local_index]),
                    "predicted_interviewer_logit_without_R": float(pqd_explanation.test_prediction[local_index]),
                    "predicted_interviewer_logit_full": float(full_explanation.test_prediction[local_index]),
                    "residual_without_R": float(residual_without_r_test[local_index]),
                    "residual_full": float(residual_full_test[local_index]),
                    "outer_train_mean_null_prediction": float(np.mean(target_train)),
                    "pqd_probability": float(probabilities["P+Q+D"][local_index]),
                    "pqd_residual_probability": float(probabilities["P+Q+D+residual_without_R"][local_index]),
                    "full_probability": float(probabilities["P+Q+R+D"][local_index]),
                    "full_residual_probability": float(probabilities["P+Q+R+D+residual_full"][local_index]),
                    "pqd_selected_C": selected_c["P+Q+D"],
                    "pqd_residual_selected_C": selected_c["P+Q+D+residual_without_R"],
                    "full_selected_C": selected_c["P+Q+R+D"],
                    "full_residual_selected_C": selected_c["P+Q+R+D+residual_full"],
                    "prediction_scope": "outer_test; source and residual construction strictly fold-contained",
                }
            )
    oof = pd.DataFrame(prediction_rows).sort_values(
        ["representation", "repeat", "outer_fold", "participant_id"], kind="stable"
    ).reset_index(drop=True)
    metrics_rows: list[dict[str, object]] = []
    for (rep_value, repeat_value), group in oof.groupby(["representation", "repeat"], sort=True):
        labels = group["label"].to_numpy(dtype=int)
        probability_columns = {
            "P+Q+D": "pqd_probability",
            "P+Q+D+residual_without_R": "pqd_residual_probability",
            "P+Q+R+D": "full_probability",
            "P+Q+R+D+residual_full": "full_residual_probability",
        }
        for model_name in R_ABSORPTION_MODELS:
            delta_from = ""
            if model_name == "P+Q+D+residual_without_R":
                delta_from = "P+Q+D"
            elif model_name == "P+Q+R+D+residual_full":
                delta_from = "P+Q+R+D"
            metrics_rows.append(
                _classification_metric_row(
                    labels,
                    group[probability_columns[model_name]].to_numpy(dtype=float),
                    representation=str(rep_value),
                    repeat=int(repeat_value),
                    model_name=model_name,
                    delta_from=delta_from,
                )
            )
        for model_name, prediction_column in (
            ("P+Q+D", "predicted_interviewer_logit_without_R"),
            ("P+Q+R+D", "predicted_interviewer_logit_full"),
        ):
            score_metrics = explanation_metrics(
                group["interviewer_logit"].to_numpy(dtype=float),
                group[prediction_column].to_numpy(dtype=float),
                group["outer_train_mean_null_prediction"].to_numpy(dtype=float),
            )
            metrics_rows.append(
                {
                    "representation": str(rep_value),
                    "repeat": int(repeat_value),
                    "model_name": model_name,
                    "metric_scope": "interviewer_score_recoverability",
                    "delta_from_model": "null",
                    **score_metrics,
                    "interpretation_scope": "exploratory_R_absorption_sensitivity_outside_primary_FDR",
                }
            )
        base_pqd = next(row for row in metrics_rows if row["model_name"] == "P+Q+D" and row["metric_scope"] == "label_prediction")
        base_full = next(row for row in metrics_rows if row["model_name"] == "P+Q+R+D" and row["metric_scope"] == "label_prediction")
        for enhanced_name, base in (
            ("P+Q+D+residual_without_R", base_pqd),
            ("P+Q+R+D+residual_full", base_full),
        ):
            enhanced = next(row for row in metrics_rows if row["model_name"] == enhanced_name and row["metric_scope"] == "label_prediction")
            metrics_rows.append(
                {
                    "representation": str(rep_value),
                    "repeat": int(repeat_value),
                    "model_name": enhanced_name,
                    "metric_scope": "paired_delta",
                    "delta_from_model": str(base["model_name"]),
                    "delta_auc": float(enhanced["roc_auc"] - base["roc_auc"]),
                    "delta_pr_auc": float(enhanced["pr_auc"] - base["pr_auc"]),
                    "delta_brier": float(base["brier"] - enhanced["brier"]),
                    "delta_log_loss": float(base["log_loss"] - enhanced["log_loss"]),
                    "interpretation_scope": "exploratory_R_absorption_sensitivity_outside_primary_FDR",
                }
            )
    return oof, pd.DataFrame(metrics_rows), pd.DataFrame(tuning_rows)


def _write_deterministic_xlsx(table: pd.DataFrame, path: Path) -> None:
    """Write a review template with normalized ZIP timestamps for rerun hashes."""

    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".xlsx.tmp")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "C5_human_validation"
    for column_index, column in enumerate(table.columns, start=1):
        sheet.cell(row=1, column=column_index, value=str(column))
    for row_index, row in enumerate(table.itertuples(index=False, name=None), start=2):
        for column_index, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column_index, value=None if value == "" else value)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.properties.creator = "DAIC-WOZ"
    workbook.properties.title = "C5 human validation candidate"
    workbook.properties.created = FIXED_XLSX_DATETIME
    workbook.properties.modified = FIXED_XLSX_DATETIME
    workbook.save(temporary)
    with zipfile.ZipFile(temporary, "r") as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
        for original, content in entries:
            info = zipfile.ZipInfo(original.filename, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            destination.writestr(info, content)
    temporary.unlink(missing_ok=True)


def write_c5_blinded_review_package(
    owner_crosswalk: pd.DataFrame,
    *,
    reviewer_dir: str | Path,
    owner_dir: str | Path,
) -> dict[str, Path]:
    """Write two blinded rater sheets and a separate restricted owner crosswalk."""

    reviewer_root = Path(reviewer_dir).resolve()
    owner_root = Path(owner_dir).resolve()
    if reviewer_root == owner_root or reviewer_root in owner_root.parents or owner_root in reviewer_root.parents:
        raise ValueError("reviewer_dir and owner_dir must be separate non-nested restricted directories")
    reviewer_root.mkdir(parents=True, exist_ok=True)
    owner_root.mkdir(parents=True, exist_ok=True)
    packet = build_c5_blinded_reviewer_packet(owner_crosswalk)
    missing_evidence = build_c5_missing_evidence_audit_form(owner_crosswalk)
    paths = {
        "reviewer_a_csv": reviewer_root / "c5_reviewer_a_blank.csv",
        "reviewer_b_csv": reviewer_root / "c5_reviewer_b_blank.csv",
        "reviewer_a_xlsx": reviewer_root / "c5_reviewer_a_blank.xlsx",
        "reviewer_b_xlsx": reviewer_root / "c5_reviewer_b_blank.xlsx",
        "missing_evidence_csv": reviewer_root / "c5_missing_evidence_audit_blank.csv",
        "owner_crosswalk_csv": owner_root / "c5_owner_crosswalk.csv",
        "owner_crosswalk_xlsx": owner_root / "c5_owner_crosswalk.xlsx",
    }
    atomic_write_csv(packet, paths["reviewer_a_csv"])
    atomic_write_csv(packet, paths["reviewer_b_csv"])
    _write_deterministic_xlsx(packet, paths["reviewer_a_xlsx"])
    _write_deterministic_xlsx(packet, paths["reviewer_b_xlsx"])
    atomic_write_csv(missing_evidence, paths["missing_evidence_csv"])
    atomic_write_csv(owner_crosswalk, paths["owner_crosswalk_csv"])
    _write_deterministic_xlsx(owner_crosswalk, paths["owner_crosswalk_xlsx"])
    return paths


def _recursive_inventory(root: Path) -> dict[str, str]:
    return {
        path.relative_to(root).as_posix(): sha256_file(path)
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def _write_preparation_outputs(
    result_dir: Path,
    *,
    project_root: Path,
    output_root: Path,
    inputs: object,
    restricted_review_dir: Path | None = None,
    restricted_owner_dir: Path | None = None,
) -> dict[str, object]:
    if (restricted_review_dir is None) != (restricted_owner_dir is None):
        raise ValueError(
            "restricted_review_dir and restricted_owner_dir must be supplied together "
            "for a blinded C5 package"
        )
    canonical = find_c5_canonical_source(project_root)
    spans = pd.read_csv(canonical)
    trace_path = output_root / "12_submission_audit/c5_traceability_audit_no_quotes.csv"
    traceability = pd.read_csv(trace_path) if trace_path.exists() else None
    evidence_counts = spans.groupby("participant_id", sort=True).size()
    roster = select_c5_validation_participants(
        inputs.structural,
        inputs.domains,
        n=24,
        evidence_counts=evidence_counts,
    )
    atomic_write_csv(roster, result_dir / "c5_human_validation_sampling_roster.csv")
    (result_dir / "c5_human_validation_field_dictionary.md").write_text(
        "# C5 human-validation field dictionary\n\n"
        "The public revision output contains only the deterministic sampling roster, "
        "field definitions, and source hashes. Reviewer packets and owner linkage "
        "are restricted artifacts and are not written here.\n\n"
        "| Field | Meaning |\n|---|---|\n"
        "| `review_case_id` | Opaque candidate-span review identifier |\n"
        "| `deidentified_quote` | Limited deterministic deidentification of the candidate quote |\n"
        "| `local_context_reference` | Opaque restricted full-text review reference |\n"
        "| `human_span_valid` | Independent human validity judgment; blank before review |\n"
        "| `human_domain` | Independent human domain judgment; blank before review |\n"
        "| `human_polarity` | Independent human polarity judgment; blank before review |\n"
        "| `rater_id` | Human coder identifier; blank before review |\n"
        "| `notes` | Human coder notes; blank before review |\n"
        "| `fulltext_review_case_id` | Opaque participant-level missing-evidence audit identifier |\n"
        "| `human_missing_evidence` | Separate full-text audit judgment, not candidate-span agreement |\n",
        encoding="utf-8",
    )
    candidate_rows = 0
    if restricted_review_dir is not None:
        owner_crosswalk = build_c5_owner_crosswalk(
            spans,
            roster,
            source_sha256=sha256_file(canonical),
            traceability=traceability,
        )
        write_c5_blinded_review_package(
            owner_crosswalk,
            reviewer_dir=restricted_review_dir,
            owner_dir=restricted_owner_dir,
        )
        candidate_rows = len(build_c5_blinded_reviewer_packet(owner_crosswalk))
    source_table = build_feature_block_source_table(inputs)
    atomic_write_csv(source_table, result_dir / "feature_block_source_table.csv")
    frozen_metrics_path = output_root / "15_interviewer_signal_explanation/final/explanation_subset_metrics.csv"
    frozen_metrics = pd.read_csv(frozen_metrics_path)
    increments = build_block_increment_contrasts(frozen_metrics)
    atomic_write_csv(increments, result_dir / "block_increment_contrasts.csv")
    frozen_source_scores_path = output_root / "15_interviewer_signal_explanation/final/source_score_crossfit.csv"
    source_performance = source_model_performance_by_repeat(frozen_source_scores_path)
    atomic_write_csv(source_performance, result_dir / "source_model_performance_by_repeat.csv")
    atomic_write_csv(
        source_model_performance_summary(source_performance),
        result_dir / "source_model_performance_summary.csv",
    )
    atomic_write_json(
        {
            "c5_canonical_source": str(C5_CANONICAL_RELATIVE).replace("\\", "/"),
            "c5_canonical_sha256": sha256_file(canonical),
            "c5_span_rows": int(len(spans)),
            "c5_span_participant_n": int(spans["participant_id"].nunique()),
            "validation_roster_n": int(len(roster)),
            "validation_candidate_rows_restricted": int(candidate_rows),
            "restricted_review_dir_used": restricted_review_dir is not None,
            "restricted_owner_dir_used": restricted_owner_dir is not None,
            "reviewer_packet_blinded": restricted_review_dir is not None,
            "owner_crosswalk_separate_and_restricted": restricted_owner_dir is not None,
            "missing_evidence_audit_scope": "separate_participant_level_full_text_review",
            "independent_human_labels_present": False,
            "c4_human_machine_agreement_reused": False,
            "c5_human_machine_agreement_status": "not_available; blinded packet and owner crosswalk prepared",
            "traceability_input_sha256": sha256_file(trace_path) if trace_path.exists() else None,
            "frozen_source_score_sha256": sha256_file(frozen_source_scores_path),
            "formal_15_unchanged": True,
        },
        result_dir / "c5_human_validation_source_hash.json",
    )
    return {
        "canonical_path": canonical,
        "canonical_sha256": sha256_file(canonical),
        "span_rows": len(spans),
        "roster_n": len(roster),
        "candidate_rows": int(candidate_rows),
        "source_score_sha256": sha256_file(frozen_source_scores_path),
        "restricted_package_written": restricted_review_dir is not None,
    }


def run_revision_analysis(
    output_root: str | Path,
    result_dir: str | Path,
    *,
    project_root: str | Path | None = None,
    code_commit: str | None = None,
    representation: str = "tfidf",
    repeat: int = 1,
    inner_folds: int = 3,
    base_seed: int = BASE_SEED,
    restricted_review_dir: str | Path | None = None,
    restricted_owner_dir: str | Path | None = None,
) -> dict[str, object]:
    """Prepare C5 material and run the locked TF-IDF repeat-1 R sensitivity."""

    output_root = Path(output_root).resolve()
    result_dir = Path(result_dir).resolve()
    if result_dir.exists():
        raise FileExistsError(f"result directory already exists: {result_dir}")
    repo_root = output_root.parent
    project = resolve_revision_project_root(output_root, project_root)
    locked_input_hashes = validate_locked_revision_inputs(output_root, project)
    inputs, membership = load_explanation_inputs(output_root, expected_n=142, expected_repeats=10, validate_hashes=True)
    result_dir.mkdir(parents=True)
    preparation = _write_preparation_outputs(
        result_dir,
        project_root=project,
        output_root=output_root,
        inputs=inputs,
        restricted_review_dir=Path(restricted_review_dir).resolve() if restricted_review_dir is not None else None,
        restricted_owner_dir=Path(restricted_owner_dir).resolve() if restricted_owner_dir is not None else None,
    )
    source_scores, source_cache, source_tuning = run_source_score_crossfit(
        inputs,
        membership,
        output_root,
        representations=(representation,),
        repeats=(int(repeat),),
        inner_folds=inner_folds,
        c_grid=DEFAULT_C_GRID,
        base_seed=base_seed,
    )
    atomic_write_csv(source_scores, result_dir / "r_absorption_source_score_crossfit.csv")
    r_oof, r_metrics, r_tuning = run_r_absorption_sensitivity(
        inputs,
        membership,
        output_root,
        representation=representation,
        repeat=repeat,
        source_cache=source_cache,
        inner_folds=inner_folds,
        alpha_grid=RIDGE_ALPHAS,
        label_c_grid=LABEL_C_GRID,
        base_seed=base_seed,
    )
    atomic_write_csv(r_oof, result_dir / "r_absorption_sensitivity_oof.csv")
    atomic_write_csv(r_metrics, result_dir / "r_absorption_sensitivity_metrics.csv")
    tuning = pd.concat(
        [source_tuning.assign(tuning_scope="source_crossfit"), r_tuning.assign(tuning_scope="R_absorption_label_models")],
        ignore_index=True,
        sort=False,
    )
    atomic_write_csv(tuning, result_dir / "r_absorption_tuning.csv")
    non_stratified_explanation_path = output_root / "15_interviewer_signal_explanation/final/fake_domain_explanation_results.csv"
    non_stratified_residual_path = output_root / "15_interviewer_signal_explanation/final/fake_domain_residual_results.csv"
    non_stratified_explanation = pd.read_csv(non_stratified_explanation_path)
    non_stratified_residual = pd.read_csv(non_stratified_residual_path)
    (
        stratified_assignments,
        stratified_explanation,
        stratified_residual,
        stratified_summary,
        stratified_explanation_oof,
        stratified_residual_oof,
    ) = run_stratified_fake_domain_controls(
        inputs,
        source_cache,
        membership,
        representation=representation,
        repeat=repeat,
        n_draws=100,
        inner_folds=inner_folds,
        alpha_grid=RIDGE_ALPHAS,
        label_c_grid=LABEL_C_GRID,
        base_seed=base_seed,
        n_bootstrap=5000,
        non_stratified_explanation_results=non_stratified_explanation,
        non_stratified_residual_results=non_stratified_residual,
    )
    atomic_write_csv(stratified_assignments, result_dir / "stratified_fake_domain_assignments.csv")
    atomic_write_csv(stratified_explanation, result_dir / "stratified_fake_domain_explanation_results.csv")
    atomic_write_csv(stratified_residual, result_dir / "stratified_fake_domain_residual_results.csv")
    atomic_write_csv(stratified_summary, result_dir / "stratified_fake_domain_summary.csv")
    atomic_write_csv(stratified_explanation_oof, result_dir / "stratified_fake_domain_explanation_oof.csv")
    atomic_write_csv(stratified_residual_oof, result_dir / "stratified_fake_domain_residual_oof.csv")
    input_hashes = {
        **locked_input_hashes,
        "splits": sha256_file(output_root / "00_splits/repeated_5fold_splits_10x5.csv"),
        "structural": sha256_file(output_root / "12_submission_audit/structural_baseline_features.csv"),
        "domain_count_D_P": sha256_file(output_root / "04_c5_controls/domain_count/input.csv"),
    }
    output_hashes = _recursive_inventory(result_dir)
    manifest = {
        "status": "complete",
        "analysis": "C5 human-validation preparation and R-block information-source/increment plus R absorption sensitivity",
        "base_formal_result_commit": BASE_FORMAL_RESULT_COMMIT,
        "analysis_code_commit": _resolve_code_commit(repo_root, code_commit),
        "cohort": {"n": 142, "validation_roster_n": int(preparation["roster_n"])},
        "representation": str(representation),
        "repeat": int(repeat),
        "feature_blocks": {block: list(columns) for block, columns in inputs.block_columns.items()},
        "r_absorption_models": list(R_ABSORPTION_MODELS),
        "source_performance": {
            "input": "analysis_v2/15_interviewer_signal_explanation/final/source_score_crossfit.csv",
            "refit_source_models": False,
            "repeat_1_bootstrap": 5000,
            "repeat_summary": "mean_sd_median_min_max; no repeat-level significance test",
        },
        "stratified_fake_D": targeted_fake_d_manifest_contract(),
        "models": {
            "explanation": {"type": "Ridge", "alpha_grid": list(RIDGE_ALPHAS), "inner_folds": int(inner_folds), "training_scaling": True},
            "label": {"type": "balanced L2 LogisticRegression", "C_grid": list(LABEL_C_GRID), "solver": "liblinear", "max_iter": 2000, "training_scaling": True},
        },
        "inference": {
            "scope": "exploratory_sensitivity_outside_frozen_four_test_family",
            "formal_hypothesis_tests": False,
        },
        "c5": {
            "canonical_sha256": preparation["canonical_sha256"],
            "independent_human_labels_present": False,
            "transcript_text_copied": False,
            "reviewer_packet_blinded_by_default": True,
            "owner_crosswalk_separate_and_restricted": True,
            "missing_evidence_scope": "separate_participant_level_full_text_review",
        },
        "input_hashes": input_hashes,
        "output_file_hashes": output_hashes,
        "output_inventory_sha256": _stable_json_hash(output_hashes),
        "completion_status": "complete",
    }
    atomic_write_json(manifest, result_dir / "run_manifest.json")
    return manifest


def verify_and_promote_revision_runs(
    module_root: str | Path,
    *,
    analysis_code_commit: str,
) -> dict[str, object]:
    """Verify two independent revision runs and recursively promote run_2."""

    module_root = Path(module_root).resolve()
    run_one = module_root / "run_1"
    run_two = module_root / "run_2"
    final = module_root / "final"
    first_inventory = _recursive_inventory(run_one)
    second_inventory = _recursive_inventory(run_two)
    if first_inventory != second_inventory:
        raise ValueError("revision run_1 and run_2 inventories are not identical")
    for run_dir in (run_one, run_two):
        manifest = json.loads((run_dir / "run_manifest.json").read_text(encoding="utf-8"))
        if manifest.get("status") != "complete":
            raise ValueError(f"revision run is not complete: {run_dir}")
    if final.exists():
        final_inventory = _recursive_inventory(final)
        if not set(final_inventory).issubset(set(second_inventory)):
            raise ValueError("existing revision final has extra files and will not be overwritten")
        if any(final_inventory[name] != second_inventory[name] for name in final_inventory):
            raise ValueError("existing revision final has hash conflicts and will not be overwritten")
        shutil.rmtree(final)
    staging = module_root / "final.staging"
    if staging.exists():
        shutil.rmtree(staging)
    shutil.copytree(run_two, staging)
    staging_inventory = _recursive_inventory(staging)
    if staging_inventory != second_inventory:
        shutil.rmtree(staging, ignore_errors=True)
        raise ValueError("revision final staging inventory does not match run_2")
    os.replace(staging, final)
    final_inventory = _recursive_inventory(final)
    if final_inventory != second_inventory:
        raise ValueError("promoted revision final inventory does not match run_2")
    verification = {
        "status": "verified",
        "run_1_inventory_sha256": _stable_json_hash(first_inventory),
        "run_2_inventory_sha256": _stable_json_hash(second_inventory),
        "run_1_run_2_identical": True,
        "promotion_status": "complete",
        "promotion_source": "run_2",
        "final_inventory_matches_run_2": True,
        "final_file_count": len(final_inventory),
        "final_inventory_sha256": _stable_json_hash(final_inventory),
        "analysis_code_commit": str(analysis_code_commit),
        "promotion_tool_commit": str(analysis_code_commit),
    }
    atomic_write_json(verification, module_root / "rerun_verification.json")
    return verification


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Prepare C5 validation data and run R-block revision sensitivity")
    parser.add_argument("--output-root", type=Path, default=Path("analysis_v2"))
    parser.add_argument("--project-root", type=Path, default=None)
    parser.add_argument("--result-dir", type=Path, default=None)
    parser.add_argument("--run-id", type=int, default=1)
    parser.add_argument("--code-commit", default=None)
    parser.add_argument("--representation", choices=("tfidf",), default="tfidf")
    parser.add_argument("--repeat", type=int, default=1)
    parser.add_argument("--restricted-review-dir", type=Path, default=None)
    parser.add_argument("--restricted-owner-dir", type=Path, default=None)
    parser.add_argument("--verify-runs", action="store_true")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    output_root = args.output_root.resolve()
    module_root = output_root / "16_interviewer_signal_revision"
    if args.verify_runs:
        analysis_commit = args.code_commit or _resolve_code_commit(output_root.parent)
        verify_and_promote_revision_runs(module_root, analysis_code_commit=analysis_commit)
        return 0
    result_dir = args.result_dir.resolve() if args.result_dir else module_root / f"run_{int(args.run_id)}"
    run_revision_analysis(
        output_root,
        result_dir,
        project_root=args.project_root,
        code_commit=args.code_commit,
        representation=args.representation,
        repeat=args.repeat,
        restricted_review_dir=args.restricted_review_dir,
        restricted_owner_dir=args.restricted_owner_dir,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
